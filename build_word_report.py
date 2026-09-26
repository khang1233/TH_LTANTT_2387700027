import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

doc = docx.Document()

# Page Margins
for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.9)
    s.right_margin = Inches(0.9)

# Logo
if os.path.exists('report_assets/hutech_logo.jpeg'):
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_after = Pt(2)
    p_logo.add_run().add_picture('report_assets/hutech_logo.jpeg', width=Inches(1.1))

# University Header
p_top = doc.add_paragraph()
p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_top.paragraph_format.space_after = Pt(2)
run = p_top.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ TP. HỒ CHÍ MINH (HUTECH)\nKHOA CÔNG NGHỆ THÔNG TIN\n")
run.font.name = 'Times New Roman'
run.font.size = Pt(12)
run.font.bold = True
run.font.color.rgb = RGBColor(0, 51, 102)

run_sub = p_top.add_run("Môn học: Thực hành Lập trình An toàn thông tin\n")
run_sub.font.name = 'Times New Roman'
run_sub.font.size = Pt(12)
run_sub.font.bold = True
run_sub.font.color.rgb = RGBColor(180, 50, 20)

p_line = doc.add_paragraph()
p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_line.paragraph_format.space_after = Pt(8)
r_line = p_line.add_run("-----------------------------------------------------------------------------------------")
r_line.font.color.rgb = RGBColor(180, 180, 180)

# Student Info Box
tbl = doc.add_table(rows=2, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for r in tbl.rows:
    for c in r.cells:
        set_cell_background(c, "F5F8FA")
        set_cell_margins(c, 70, 70, 120, 120)

def add_box_field(cell, label, val):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(label)
    r1.font.name = 'Times New Roman'
    r1.font.bold = True
    r2 = p.add_run(val)
    r2.font.name = 'Times New Roman'

add_box_field(tbl.cell(0, 0), "Họ và Tên: ", "Võ Duy Khang")
add_box_field(tbl.cell(0, 1), "MSSV: ", "2387700027")
add_box_field(tbl.cell(1, 0), "Lớp: ", "23DATA1 / ATTT")
add_box_field(tbl.cell(1, 1), "GitHub: ", "khang1233/TH_LTANTT_2387700027")

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(14)
p_title.paragraph_format.space_after = Pt(14)
run_t = p_title.add_run("BÁO CÁO THỰC HÀNH - BÀI 1: LẬP TRÌNH CƠ BẢN VỚI PYTHON")
run_t.font.name = 'Times New Roman'
run_t.font.size = Pt(14)
run_t.font.bold = True
run_t.font.color.rgb = RGBColor(0, 51, 102)

# List of figures
figures = [
    ("Hình 1: Kết quả chạy chương trình hello.py", "report_assets/hinh1_hello.png"),
    ("Hình 2: Kết quả thực thi ex02_01.py (Nhập họ tên và tuổi)", "report_assets/hinh2_ex02_01.png"),
    ("Hình 3: Kết quả thực thi ex02_02.py (Tính diện tích hình tròn)", "report_assets/hinh3_ex02_02.png"),
    ("Hình 4: Kết quả thực thi ex02_03.py (Kiểm tra số chẵn / số lẻ)", "report_assets/hinh4_ex02_03.png"),
    ("Hình 5: Kết quả thực thi ex02_04.py (Dãy số chia hết cho 7 không chia hết cho 5)", "report_assets/hinh5_ex02_04.png"),
    ("Hình 6: Kết quả thực thi ex02_05.py (Tính tiền lương nhân viên)", "report_assets/hinh6_ex02_05.png"),
    ("Hình 7: Kết quả thực thi ex02_06.py (Tạo mảng 2 chiều X x Y)", "report_assets/hinh7_ex02_06.png"),
    ("Hình 8: Kết quả thực thi ex02_07.py (Chuyển chuỗi thành chữ in hoa)", "report_assets/hinh8_ex02_07.png"),
    ("Hình 9: Kết quả thực thi ex02_08.py (Lọc số nhị phân chia hết cho 5)", "report_assets/hinh9_ex02_08.png"),
    ("Hình 10: Kết quả thực thi ex02_09.py (Hàm kiểm tra số nguyên tố)", "report_assets/hinh10_ex02_09.png"),
    ("Hình 11: Kết quả thực thi ex02_10.py (Hàm đảo ngược chuỗi)", "report_assets/hinh11_ex02_10.png"),
    ("Hình 12: Kết quả thực thi ex03_01.py (Tổng các số chẵn trong List)", "report_assets/hinh12_ex03_01.png"),
    ("Hình 13: Kết quả thực thi ex03_02.py (Đảo ngược List)", "report_assets/hinh13_ex03_02.png"),
    ("Hình 14: Kết quả thực thi ex03_03.py (Tạo Tuple từ List)", "report_assets/hinh14_ex03_03.png"),
    ("Hình 15: Kết quả thực thi ex03_04.py (Truy cập phần tử đầu và cuối Tuple)", "report_assets/hinh15_ex03_04.png"),
    ("Hình 16: Kết quả thực thi ex03_05.py (Đếm tần suất từ vào Dictionary)", "report_assets/hinh16_ex03_05.png"),
    ("Hình 17: Kết quả thực thi ex03_06.py (Xóa phần tử Dictionary theo khóa)", "report_assets/hinh17_ex03_06.png"),
    ("Hình 18: Kết quả thực thi Main.py (Ứng dụng Quản lý Sinh viên OOP)", "report_assets/hinh18_ex04_main.png"),
    ("Hình 19: Trạng thái Git Status và Push toàn bộ mã nguồn lên GitHub", "report_assets/hinh19_git_push.png"),
]

for title, img_path in figures:
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_before = Pt(10)
    p_cap.paragraph_format.space_after = Pt(4)
    p_cap.paragraph_format.keep_with_next = True
    r = p_cap.add_run(title)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(20, 40, 80)
    
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(10)
        p_img.add_run().add_picture(img_path, width=Inches(5.9))

output_path = "BaoCao_Lab01.docx"
doc.save(output_path)
print(f"Updated {output_path} successfully!")
try:
    doc.save("BaoCao_ThucHanh_Lab01.docx")
except Exception:
    pass
