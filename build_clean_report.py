import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()

# Margins 2cm
for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)

# Logo HUTECH
if os.path.exists('report_assets/hutech_logo.jpeg'):
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_after = Pt(2)
    p_logo.add_run().add_picture('report_assets/hutech_logo.jpeg', width=Inches(1.0))

# Header don gian, chu den
p_top = doc.add_paragraph()
p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_top.paragraph_format.space_after = Pt(2)
r = p_top.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ TP. HỒ CHÍ MINH (HUTECH)\nKHOA CÔNG NGHỆ THÔNG TIN\n")
r.font.name = 'Times New Roman'
r.font.size = Pt(12)
r.font.bold = True
r.font.color.rgb = RGBColor(0, 0, 0)

r2 = p_top.add_run("Môn học: Thực hành Lập trình An toàn thông tin\n")
r2.font.name = 'Times New Roman'
r2.font.size = Pt(12)
r2.font.bold = True
r2.font.color.rgb = RGBColor(0, 0, 0)

# Thong tin sinh vien
p_info = doc.add_paragraph()
p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_info.paragraph_format.space_after = Pt(6)
r3 = p_info.add_run("Họ và Tên: Trần Minh Khang   -   MSSV: 2387700027   -   Lớp: 23DATA1 / ATTT\nGitHub: https://github.com/khang1233/TH_LTANTT_2387700027")
r3.font.name = 'Times New Roman'
r3.font.size = Pt(11)
r3.font.color.rgb = RGBColor(0, 0, 0)

# Tieu de bao cao
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(8)
p_title.paragraph_format.space_after = Pt(14)
r_title = p_title.add_run("BÁO CÁO THỰC HÀNH - BÀI 1: LẬP TRÌNH CƠ BẢN VỚI NGÔN NGỮ PYTHON")
r_title.font.name = 'Times New Roman'
r_title.font.size = Pt(13)
r_title.font.bold = True
r_title.font.color.rgb = RGBColor(0, 0, 0)

# Danh sach cac hinh: khong co code, chi tieu de den + anh
items = [
    ("Hình 1: Kết quả chạy chương trình hello.py", "report_assets/hinh1_hello.png"),
    ("Hình 2: ex02_01.py - Nhập họ tên và tuổi", "report_assets/hinh2_ex02_01.png"),
    ("Hình 3: ex02_02.py - Tính diện tích hình tròn (Pi = 3.14)", "report_assets/hinh3_ex02_02.png"),
    ("Hình 4: ex02_03.py - Kiểm tra số chẵn / số lẻ", "report_assets/hinh4_ex02_03.png"),
    ("Hình 5: ex02_04.py - Tìm số chia hết cho 7 không chia hết cho 5 trong đoạn [2000, 3200]", "report_assets/hinh5_ex02_04.png"),
    ("Hình 6: ex02_05.py - Tính tiền lương thực nhận của nhân viên", "report_assets/hinh6_ex02_05.png"),
    ("Hình 7: ex02_06.py - Tạo mảng 2 chiều X x Y với giá trị phần tử i * j", "report_assets/hinh7_ex02_06.png"),
    ("Hình 8: ex02_07.py - Chuyển đổi các dòng văn bản thành chữ in hoa", "report_assets/hinh8_ex02_07.png"),
    ("Hình 9: ex02_08.py - Lọc các số nhị phân 4 chữ số chia hết cho 5", "report_assets/hinh9_ex02_08.png"),
    ("Hình 10: ex02_09.py - Hàm kiểm tra số nguyên tố", "report_assets/hinh10_ex02_09.png"),
    ("Hình 11: ex02_10.py - Hàm đảo ngược chuỗi ký tự", "report_assets/hinh11_ex02_10.png"),
    ("Hình 12: ex03_01.py - Tính tổng các số chẵn trong List", "report_assets/hinh12_ex03_01.png"),
    ("Hình 13: ex03_02.py - Đảo ngược vị trí các phần tử trong List", "report_assets/hinh13_ex03_02.png"),
    ("Hình 14: ex03_03.py - Tạo một Tuple từ một List nhập vào", "report_assets/hinh14_ex03_03.png"),
    ("Hình 15: ex03_04.py - Truy cập phần tử đầu tiên và cuối cùng trong Tuple", "report_assets/hinh15_ex03_04.png"),
    ("Hình 16: ex03_05.py - Đếm số lần xuất hiện của từ vào Dictionary", "report_assets/hinh16_ex03_05.png"),
    ("Hình 17: ex03_06.py - Xóa phần tử khỏi Dictionary theo khóa", "report_assets/hinh17_ex03_06.png"),
    ("Hình 18: Main.py - Kết quả thực thi ứng dụng Quản lý Sinh viên OOP", "report_assets/hinh18_ex04_main.png"),
    ("Hình 19: Trạng thái Git Status và Push lên GitHub", "report_assets/hinh19_git_push.png")
]

for title, img_path in items:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run_t = p.add_run(title)
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(11)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0, 0, 0)
    
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(12)
        p_img.add_run().add_picture(img_path, width=Inches(5.9))

for fname in ["BaoCao_Lab01_Final.docx", "BaoCao_Lab01.docx", "BaoCao_Lab01_HoanChinh.docx"]:
    try:
        doc.save(fname)
        print(f"Saved: {fname}")
    except Exception as e:
        print(f"Skipped {fname} (likely open in Word): {e}")
